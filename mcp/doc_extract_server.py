#!/usr/bin/env python3
"""MCP server: doc_extract — convert PDF/DOCX/HTML/XLSX to markdown text.

Lightweight alternative to docling: uses pypdfium2 (PDF), python-docx (DOCX),
markdownify (HTML), openpyxl (XLSX). No torch/cv2 dependency.
"""
import asyncio
import json
import os
import tempfile
import urllib.request
from pathlib import Path
from typing import Any

from mcp.server import Server
from mcp.server.stdio import stdio_server
from mcp.types import TextContent, Tool

server = Server("doc-extract")

MAX_BYTES = int(os.environ.get("DOC_EXTRACT_MAX_BYTES", str(20 * 1024 * 1024)))  # 20MB


def _extract_pdf(path: str) -> str:
    import pypdfium2 as pdfium
    pdf = pdfium.PdfDocument(path)
    pages = []
    for i, page in enumerate(pdf):
        textpage = page.get_textpage()
        text = textpage.get_text_range()
        if text.strip():
            pages.append(f"## Page {i+1}\n\n{text}")
        textpage.close()
        page.close()
    pdf.close()
    return "\n\n".join(pages) if pages else "(no text extracted)"


def _extract_docx(path: str) -> str:
    from docx import Document
    doc = Document(path)
    parts = []
    for para in doc.paragraphs:
        if para.text.strip():
            style = para.style.name.lower() if para.style else ""
            if "heading 1" in style:
                parts.append(f"# {para.text}")
            elif "heading 2" in style:
                parts.append(f"## {para.text}")
            elif "heading 3" in style:
                parts.append(f"### {para.text}")
            else:
                parts.append(para.text)
    for table in doc.tables:
        for row in table.rows:
            parts.append("| " + " | ".join(cell.text.replace("\n", " ") for cell in row.cells) + " |")
    return "\n\n".join(parts) if parts else "(empty document)"


def _extract_doc(path: str) -> str:
    """.doc (legacy binary) — через Word COM Content.Text (win32com).

    НЕ SaveAs: 'Ошибка метода'/method error (нестабилен в COM из PS 5.1).
    """
    import pythoncom
    import win32com.client

    pythoncom.CoInitialize()
    word = None
    try:
        word = win32com.client.DispatchEx("Word.Application")
        word.Visible = False
        word.DisplayAlerts = 0
        doc = word.Documents.Open(os.path.abspath(path), False, True)
        try:
            text = doc.Content.Text
        finally:
            doc.Close(False)
        return text
    finally:
        try:
            if word is not None:
                word.Quit()
        except Exception:
            pass
        pythoncom.CoUninitialize()


def _extract_html(path: str) -> str:
    from markdownify import markdownify as md
    with open(path, encoding="utf-8", errors="replace") as f:
        html = f.read()
    return md(html, heading_style="ATX")


def _extract_xlsx(path: str) -> str:
    from openpyxl import load_workbook
    wb = load_workbook(path, read_only=True, data_only=True)
    parts = []
    for ws in wb.worksheets:
        parts.append(f"## Sheet: {ws.title}\n")
        for row in ws.iter_rows(max_row=200, values_only=True):
            if any(c is not None for c in row):
                parts.append("| " + " | ".join(str(c) if c is not None else "" for c in row) + " |")
        parts.append("")
    wb.close()
    return "\n".join(parts) if parts else "(empty spreadsheet)"


def _extract_plain(path: str) -> str:
    return Path(path).read_text(encoding="utf-8", errors="replace")


EXTRACTORS = {
    ".pdf": _extract_pdf,
    ".docx": _extract_docx,
    ".doc": _extract_doc,
    ".html": _extract_html,
    ".htm": _extract_html,
    ".xlsx": _extract_xlsx,
    ".txt": _extract_plain,
    ".md": _extract_plain,
    ".csv": _extract_plain,
    ".json": _extract_plain,
}


def _do_extract(path: str) -> dict:
    p = Path(path)
    if not p.exists():
        return {"ok": False, "error": f"file not found: {path}"}
    size = p.stat().st_size
    if size > MAX_BYTES:
        return {"ok": False, "error": f"file too large: {size} bytes (max {MAX_BYTES})"}
    ext = p.suffix.lower()
    extractor = EXTRACTORS.get(ext)
    if not extractor:
        return {"ok": False, "error": f"unsupported format: {ext}", "supported": sorted(EXTRACTORS.keys())}
    text = extractor(str(p))
    return {"ok": True, "file": str(p), "size_bytes": size, "format": ext, "chars": len(text), "text": text}


@server.list_tools()
async def list_tools() -> list[Tool]:
    return [
        Tool(
            name="extract_document",
            description="Extract text from a local file (PDF/DOCX/DOC/HTML/XLSX/TXT/MD/CSV) as markdown. "
                        ".doc via Word COM Content.Text. Use for research-academic: parse downloaded papers, reports, documents.",
            inputSchema={
                "type": "object",
                "properties": {
                    "path": {"type": "string", "description": "Absolute path to the local file."},
                },
                "required": ["path"],
            },
        )
    ]


@server.call_tool()
async def call_tool(name: str, arguments: Any) -> list[TextContent]:
    if name != "extract_document":
        return [TextContent(type="text", text=json.dumps({"ok": False, "error": f"Unknown tool: {name}"}))]
    if isinstance(arguments, str):
        arguments = json.loads(arguments)
    args = arguments or {}
    path = str(args.get("path", "")).strip()
    if not path:
        return [TextContent(type="text", text=json.dumps({"ok": False, "error": "path is required"}))]
    loop = asyncio.get_event_loop()
    try:
        result = await loop.run_in_executor(None, _do_extract, path)
    except Exception as exc:
        result = {"ok": False, "error": str(exc)}
    return [TextContent(type="text", text=json.dumps(result, ensure_ascii=False, indent=2))]


async def main():
    async with stdio_server() as (read_stream, write_stream):
        await server.run(read_stream, write_stream, server.create_initialization_options())


if __name__ == "__main__":
    asyncio.run(main())