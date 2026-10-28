"""Bounded local extraction. Never executes corpus code or downloads dependencies."""
from dataclasses import dataclass
from html import unescape
from pathlib import Path
import re
import zipfile
import importlib.metadata
from urllib.parse import parse_qsl, urlencode, urlsplit, urlunsplit


TEXT_SUFFIXES = {
    ".md", ".txt", ".rst", ".json", ".jsonl", ".yaml", ".yml", ".csv",
    ".tsv", ".py", ".js", ".ts", ".tsx", ".jsx", ".ps1", ".sh", ".toml",
    ".xml", ".html", ".htm", ".tex", ".bib", ".ini", ".log",
}
SUPPORTED = TEXT_SUFFIXES | {".pdf", ".docx", ".xlsx", ".zip"}


def parser_versions():
    versions = []
    for name in ("PyMuPDF", "python-docx", "openpyxl"):
        try:
            versions.append(name + "=" + importlib.metadata.version(name))
        except importlib.metadata.PackageNotFoundError:
            versions.append(name + "=absent")
    return ";".join(versions)


@dataclass
class Extraction:
    text: str = ""
    status: str = "extracted"
    detail: str = ""


SECRET_FIELD = r"(?:api[_-]?key|password|secret|(?:(?:access|refresh|auth)[_-]?)?token|client[_-]?secret|authorization|private[_-]?key)"


def redact_url(value: str) -> str:
    try:
        parsed = urlsplit(value)
        if parsed.scheme.lower() not in {"https", "http"} or not parsed.hostname:
            return value
        host = parsed.netloc.rsplit("@", 1)[-1]
        query = [(key, "REDACTED" if re.search(r"(?i)token|key|secret|password|signature|authorization", key) else val)
                 for key, val in parse_qsl(parsed.query, keep_blank_values=True)]
        fragment = re.sub(r"(?i)((?:token|key|secret|password|signature)=)[^&]*", r"\1REDACTED", parsed.fragment)
        return urlunsplit((parsed.scheme, host, parsed.path, urlencode(query), fragment))
    except ValueError:
        return "<redacted-invalid-url>"


def redact(text: str) -> str:
    text = re.sub(r"(?i)(bearer\s+)[A-Za-z0-9._~+/-]{12,}", r"\1<redacted>", text)
    text = re.sub(r"\bsk-[A-Za-z0-9_-]{16,}\b", "<redacted-key>", text)
    text = re.sub(rf'''(?i)(["']{SECRET_FIELD}["']\s*:\s*)(["'])(?:\\.|(?!\2).)*?\2''', r'\1"<redacted>"', text)
    text = re.sub(r"https?://[^\s<>\"'`]+", lambda match: redact_url(match[0]), text)
    return re.sub(rf"(?im)^(\s*[\w.-]*{SECRET_FIELD}[\w.-]*\s*[:=]\s*).+$", r"\1<redacted>", text)


def safe_data(value):
    """Sanitize structured worker/legacy outputs without corrupting JSON syntax."""
    if isinstance(value, dict):
        return {key: "<redacted>" if re.fullmatch(SECRET_FIELD, str(key), re.I) else safe_data(item)
                for key, item in value.items()}
    if isinstance(value, list):
        return [safe_data(item) for item in value]
    return redact(value) if isinstance(value, str) else value


def extract(path: Path, max_chars: int = 250000, max_pages: int = 200) -> Extraction:
    suffix = path.suffix.lower()
    if suffix not in SUPPORTED:
        return Extraction(status="unsupported", detail=f"unsupported format: {suffix}")
    partial = False
    try:
        if suffix in TEXT_SUFFIXES:
            data = path.read_bytes()
            try:
                text = data.decode("utf-8-sig")
            except UnicodeDecodeError:
                text = data.decode("cp1251", errors="replace")
            if "\x00" in text:
                return Extraction(status="failed", detail="binary data in text format")
            if suffix in {".html", ".htm"}:
                text = re.sub(r"(?is)<(script|style|nav)\b.*?</\1\s*>", " ", text)
                text = unescape(re.sub(r"(?s)<[^>]+>", " ", text))
        elif suffix == ".pdf":
            import fitz
            with fitz.open(path) as doc:
                partial = len(doc) > max_pages
                pieces = []
                size = 0
                for i in range(min(len(doc), max_pages)):
                    piece = f"\n## Page {i + 1}\n{doc[i].get_text()}"
                    pieces.append(piece)
                    size += len(piece)
                    if size > max_chars:
                        partial = True
                        break
                text = "\n".join(pieces)
                if not any(re.search(r"\w", piece.split("\n", 2)[-1]) for piece in pieces):
                    return Extraction(status="needs_ocr", detail="no extractable PDF text")
        elif suffix == ".docx":
            from docx import Document
            with zipfile.ZipFile(path) as archive:
                if sum(x.file_size for x in archive.infolist()) > 128 * 1024 * 1024:
                    return Extraction(status="failed", detail="DOCX expanded size cap")
            doc = Document(path)
            text = "\n".join([p.text for p in doc.paragraphs] + [" | ".join(c.text for c in r.cells) for t in doc.tables for r in t.rows])
        elif suffix == ".xlsx":
            from openpyxl import load_workbook
            book = load_workbook(path, read_only=True, data_only=True)
            try:
                pieces = []
                partial = len(book.worksheets) > 20
                for sheet in book.worksheets[:20]:
                    pieces.append(f"## Sheet {sheet.title}")
                    partial |= (sheet.max_row or 0) > 2000 or (sheet.max_column or 0) > 50
                    for row in sheet.iter_rows(max_row=2000, max_col=50, values_only=True):
                        if any(x is not None for x in row):
                            pieces.append(" | ".join(str(x) if x is not None else "" for x in row))
                text = "\n".join(pieces)
            finally:
                book.close()
        else:
            with zipfile.ZipFile(path) as archive:
                entries = archive.infolist()
                text = "Archive members (contents NOT extracted):\n" + "\n".join(x.filename for x in entries[:2000])
            return Extraction(redact(text)[:max_chars], "partial", "archive inventory only")
        partial |= len(text) > max_chars
        return Extraction(redact(text)[:max_chars], "partial" if partial else "extracted", "extraction limit" if partial else "")
    except Exception as exc:
        return Extraction(status="failed", detail=f"{type(exc).__name__}: {str(exc)[:200]}")
