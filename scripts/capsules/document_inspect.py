#!/usr/bin/env python3
from __future__ import annotations
import argparse, csv, hashlib, json, os
from pathlib import Path

def sha256_file(p:Path)->str:
    h=hashlib.sha256()
    with p.open('rb') as f:
        for b in iter(lambda:f.read(1024*1024),b''): h.update(b)
    return h.hexdigest()

def inspect_text(p:Path):
    text=p.read_text(encoding='utf-8',errors='replace')
    return {'kind':'text','segments':[{'locator':'line:1','text':text[:200000]}], 'warnings':['truncated'] if len(text)>200000 else []}

def inspect_pdf(p:Path):
    try:
        import pymupdf as fitz
    except ImportError:
        import fitz
    doc=fitz.open(p)
    seg=[]; warnings=[]
    for i,page in enumerate(doc):
        txt=page.get_text('text') or ''
        if not txt.strip(): warnings.append(f'page:{i+1}:no_text_layer')
        seg.append({'locator':f'pdf_page:{i+1}','text':txt[:100000],'text_chars':len(txt)})
    return {'kind':'pdf','page_count':len(doc),'segments':seg,'warnings':warnings}

def inspect_docx(p:Path):
    from docx import Document
    d=Document(p); seg=[]
    for i,par in enumerate(d.paragraphs,1):
        if par.text.strip(): seg.append({'locator':f'paragraph:{i}','text':par.text})
    tables=[]
    for ti,t in enumerate(d.tables,1):
        rows=[]
        for ri,row in enumerate(t.rows,1):
            vals=[cell.text for cell in row.cells]; rows.append(vals)
        tables.append({'locator':f'table:{ti}','rows':rows})
    return {'kind':'docx','segments':seg,'tables':tables,'warnings':[]}

def inspect_xlsx(p:Path):
    import openpyxl
    wb_formula=openpyxl.load_workbook(p,data_only=False,read_only=False)
    wb_values=openpyxl.load_workbook(p,data_only=True,read_only=False)
    sheets=[]; warnings=[]
    for ws in wb_formula.worksheets:
        wsv=wb_values[ws.title]; cells=[]
        for row in ws.iter_rows():
            for c in row:
                if c.value is None: continue
                cv=wsv[c.coordinate].value
                cells.append({'cell':c.coordinate,'formula':c.value if isinstance(c.value,str) and c.value.startswith('=') else None,'value':cv if isinstance(c.value,str) and c.value.startswith('=') else c.value,'cached_value':cv,'number_format':c.number_format,'hidden_row':bool(ws.row_dimensions[c.row].hidden),'hidden_col':bool(ws.column_dimensions[c.column_letter].hidden)})
        sheets.append({'sheet':ws.title,'hidden':ws.sheet_state!='visible','merged_ranges':[str(x) for x in ws.merged_cells.ranges],'cells':cells})
    return {'kind':'xlsx','sheets':sheets,'warnings':warnings}

def inspect_tiff(p:Path):
    import tifffile
    frames=[]
    with tifffile.TiffFile(p) as tf:
        for i,page in enumerate(tf.pages):
            tags={}
            for t in page.tags.values():
                try:
                    v=t.value
                    if isinstance(v,(bytes,bytearray)): v=f'<bytes:{len(v)}>'
                    elif not isinstance(v,(str,int,float,bool,list,tuple,dict,type(None))): v=str(v)
                    tags[str(t.name)]=v
                except Exception: pass
            frames.append({'frame':i,'shape':list(page.shape),'dtype':str(page.dtype),'tags':tags})
    return {'kind':'tiff','frames':frames,'warnings':['vendor_private_tags_not_semantically_interpreted']}

def inspect_csv(p:Path):
    with p.open('r',encoding='utf-8-sig',errors='replace',newline='') as f: rows=list(csv.reader(f))
    return {'kind':'csv','rows':rows[:10000],'warnings':['truncated'] if len(rows)>10000 else []}

def main()->int:
    ap=argparse.ArgumentParser(); ap.add_argument('path'); args=ap.parse_args(); p=Path(args.path).expanduser().resolve()
    if not p.exists() or not p.is_file():
        print(json.dumps({'ok':False,'error':'file_not_found','path':str(p)},ensure_ascii=False,indent=2)); return 2
    ext=p.suffix.lower()
    try:
        if ext in {'.txt','.md','.json','.yaml','.yml','.py','.log'}: body=inspect_text(p)
        elif ext=='.pdf': body=inspect_pdf(p)
        elif ext=='.docx': body=inspect_docx(p)
        elif ext=='.xlsx': body=inspect_xlsx(p)
        elif ext in {'.tif','.tiff'}: body=inspect_tiff(p)
        elif ext=='.csv': body=inspect_csv(p)
        elif ext in {'.doc','.xls'}:
            body={'kind':'legacy_office','warnings':['legacy_office_requires_conversion_provider'],'segments':[]}
        else: body={'kind':'unsupported','warnings':[f'unsupported_extension:{ext}']}
        out={'ok':body.get('kind')!='unsupported','schema':'document_inspection.v1','path':str(p),'sha256':sha256_file(p),'size_bytes':p.stat().st_size,**body}
    except Exception as e:
        out={'ok':False,'schema':'document_inspection.v1','path':str(p),'sha256':sha256_file(p),'error':f'{e.__class__.__name__}: {e}'}
    print(json.dumps(out,ensure_ascii=False,indent=2,default=str)); return 0 if out.get('ok') else 3
if __name__=='__main__': raise SystemExit(main())
