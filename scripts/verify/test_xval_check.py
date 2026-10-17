#!/usr/bin/env python3
"""pytest для xval_check.py — 5 кейсов исследования P1a (TD-FACTORY-1).

Кейсы:
1. баг X==1..N → FAIL (issue enum_reset_x)
2. правильная linspace-сетка → PASS
3. немонотонность (перестановка точки) → FAIL (non_monotonic)
4. len(X)!=len(Y) → FAIL (len_mismatch)
5. Y — живые формулы без кэша (data_only читает None) → PASS + warning len_formula_no_cache

Книги строятся в памяти (openpyxl), тесты без round-trip вызывают check_xval(wb)
на открытой книге. Кейс 5 требует round-trip через файл (data_only=True), т.к.
в памяти формула возвращает строку '=G*E', а кэша нет.
"""

import os
import tempfile

import numpy as np
import pytest

from openpyxl import Workbook, load_workbook
from openpyxl.chart import ScatterChart, Reference
from openpyxl.chart.data_source import AxDataSource, NumRef, NumDataSource
from openpyxl.chart.series import XYSeries

import importlib.util
from pathlib import Path

_HERE = Path(__file__).resolve().parent
_spec = importlib.util.spec_from_file_location("verify_xval_check_mod", _HERE / "xval_check.py")
assert _spec and _spec.loader
_xval_mod = importlib.util.module_from_spec(_spec)
_spec.loader.exec_module(_xval_mod)
check_xval = _xval_mod.check_xval


def _make_scatter_book(xs, ys) -> Workbook:
    """ScatterChart, 1 серия: X = xs (столбец A), Y = ys (столбец B)."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Лист1"
    n = max(len(xs), len(ys))
    for i in range(n):
        if i < len(xs):
            ws.cell(row=1 + i, column=1, value=xs[i])
        if i < len(ys):
            ws.cell(row=1 + i, column=2, value=ys[i])
    chart = ScatterChart()
    xref = Reference(ws, min_col=1, min_row=1, max_row=max(1, len(xs)))
    yref = Reference(ws, min_col=2, min_row=1, max_row=max(1, len(ys)))
    ser = XYSeries(
        idx=0, order=0,
        xVal=AxDataSource(numRef=NumRef(f=str(xref))),
        yVal=NumDataSource(numRef=NumRef(f=str(yref))),
    )
    chart.series.append(ser)
    ws.add_chart(chart, "D1")
    return wb


def _save_roundtrip(wb) -> str:
    fd, path = tempfile.mkstemp(suffix=".xlsx")
    os.close(fd)
    wb.save(path)
    return path


# ---------------------------------------------------------------- кейсы

def test_enum_reset_x_bug_fails():
    """Кейс 1: X==1..N — сигнатура Excel-бага → FAIL enum_reset_x."""
    wb = _make_scatter_book([1, 2, 3, 4, 5, 6], [1.0, 4.0, 9.0, 16.0, 25.0, 36.0])
    result = check_xval(wb)
    assert result["ok"] is False
    issues = {i["issue"] for i in result["issues"]}
    assert "enum_reset_x" in issues


def test_correct_grid_passes():
    """Кейс 2: правильная linspace-сетка → PASS (ok True, issues пусто)."""
    xs = np.linspace(0.02, 3.1, 200)
    ys = np.linspace(0.02, 3.1, 200) ** 2
    wb = _make_scatter_book(xs.tolist(), ys.tolist())
    result = check_xval(wb)
    assert result["ok"] is True
    assert result["issues"] == []
    assert result["series_count"] == 1


def test_non_monotonic_fails():
    """Кейс 3: немонотонность (одна точка заменена на предыдущую) → FAIL non_monotonic."""
    xs = np.linspace(0.02, 3.1, 200).tolist()
    # заменить точку на соседнюю — строгая монотонность ломается
    xs[50] = xs[49]
    ys = np.linspace(0.02, 3.1, 200).tolist()
    wb = _make_scatter_book(xs, ys)
    result = check_xval(wb)
    assert result["ok"] is False
    issues = {i["issue"] for i in result["issues"]}
    assert "non_monotonic" in issues


def test_len_mismatch_fails():
    """Кейс 4: len(X)!=len(Y) → FAIL len_mismatch."""
    wb = _make_scatter_book([0.0, 1.0, 2.0, 3.0, 4.0, 5.0], [0.0, 1.0, 4.0, 9.0, 16.0])
    result = check_xval(wb)
    assert result["ok"] is False
    issues = {i["issue"] for i in result["issues"]}
    assert "len_mismatch" in issues


def test_live_formula_y_no_cache_passes_with_warning():
    """Кейс 5: Y — живые формулы без кэша (lenY=0) → PASS + warning len_formula_no_cache.

    Строится книга с Y = '=G..*E..' и сохраняется на диск: data_only=True
    читает кэш, которого нет → Y=[]; по размеру ссылки проверка не фейлит.
    """
    wb = Workbook()
    ws = wb.active
    ws.title = "Лист1"
    xs = np.linspace(0.02, 3.1, 6)
    for i in range(6):
        ws.cell(row=1 + i, column=1, value=float(xs[i]))
        ws.cell(row=1 + i, column=2, value=f"=G{1 + i}*E{1 + i}")  # живая формула без кэша
    chart = ScatterChart()
    xref = Reference(ws, min_col=1, min_row=1, max_row=6)
    yref = Reference(ws, min_col=2, min_row=1, max_row=6)
    ser = XYSeries(
        idx=0, order=0,
        xVal=AxDataSource(numRef=NumRef(f=str(xref))),
        yVal=NumDataSource(numRef=NumRef(f=str(yref))),
    )
    chart.series.append(ser)
    ws.add_chart(chart, "D1")
    path = _save_roundtrip(wb)
    try:
        result = check_xval(path)  # data_only=True
        assert result["ok"] is True
        warnings = {w["issue"] for w in result["warnings"]}
        assert "len_formula_no_cache" in warnings
    finally:
        os.remove(path)


# ---------------------------------------------------------------- вспомогательные

def test_sheet_not_found():
    """Запрошенного листа нет → FAIL sheet_not_found."""
    wb = _make_scatter_book([0.0, 1.0], [0.0, 1.0])
    result = check_xval(wb, sheet_name="Nope")
    assert result["ok"] is False
    assert result["issues"][0]["issue"] == "sheet_not_found"


def test_no_xval_issue():
    """Серия без xVal и без cat → FAIL no_xval."""
    wb = Workbook()
    ws = wb.active
    from openpyxl.chart import PieChart
    chart = PieChart()
    chart.add_data(Reference(ws, min_col=1, min_row=1, max_row=3), titles_from_data=False)
    ws["A1"], ws["A2"], ws["A3"] = 1, 2, 3
    ws.add_chart(chart, "C1")
    result = check_xval(wb)
    assert result["ok"] is False
    assert any(i["issue"] == "no_xval" for i in result["issues"])


def test_grid_filter_skips_calibration_points():
    """Неравномерный ряд (точки калибровки) не фейлится при grid_filter."""
    xs = [0.026, 0.0919, 0.4848, 0.9905, 2.2233, 3.1443]  # реальные уровни, НЕ linspace
    ys = [v * 2 for v in xs]
    wb = _make_scatter_book(xs, ys)
    # без фильтра — grid_mismatch (неравномерный ряд)
    assert not check_xval(wb)["ok"]
    # с автофильтром — только uniform-ряды сверяются → PASS
    result = check_xval(wb, grid_filter=_xval_mod._auto_grid_filter)
    assert result["ok"] is True


def test_allow_singletons_skips_marker():
    """Одиночный маркер (len(X)==1) с allow_singletons=True → PASS + warning."""
    wb = _make_scatter_book([0.9966], [12.5])
    assert not check_xval(wb)["ok"]  # degenerate issue
    result = check_xval(wb, allow_singletons=True)
    assert result["ok"] is True
    assert any(w["issue"] == "singleton_marker" for w in result["warnings"])


def test_check_xval_returns_contract_shape():
    """Контракт CheckResult: ключи ok/series_count/issues/xvals/warnings."""
    wb = _make_scatter_book([0.0, 1.0, 2.0], [0.0, 1.0, 4.0])
    result = check_xval(wb)
    assert set(result.keys()) == {"ok", "series_count", "issues", "xvals", "warnings"}
    for i in result["issues"]:
        assert {"series_idx", "issue", "detail"} <= set(i.keys())