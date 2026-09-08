import json, os, subprocess, sys, tempfile, unittest
from pathlib import Path
ROOT=Path(os.environ.get('OPENCODE_HARNESS_ROOT','')).resolve()

def run(*args):
    cp=subprocess.run([sys.executable,*map(str,args)],cwd=ROOT,text=True,capture_output=True)
    try: data=json.loads(cp.stdout) if cp.stdout.strip().startswith('{') else None
    except Exception: data=None
    return cp,data

class OverlayTests(unittest.TestCase):
    def test_compile_check(self):
        cp,d=run(ROOT/'scripts/router/compile_capability_runtime.py','--check'); self.assertEqual(cp.returncode,0,cp.stderr+cp.stdout); self.assertTrue(d['ok'])

    def test_preflight_distinguishes_degraded(self):
        cp,d=run(ROOT/'scripts/router/capability_preflight.py','--no-network'); self.assertEqual(cp.returncode,0,cp.stderr); self.assertEqual(d['capabilities']['web.discovery']['status'],'degraded'); self.assertTrue(d['capabilities']['science.numeric']['available']); self.assertTrue(d['capabilities']['job.checkpoint']['available'])

    def test_research_bundle_any_of_fallback(self):
        cp,pf=run(ROOT/'scripts/router/capability_preflight.py','--no-network'); self.assertEqual(cp.returncode,0)
        # deterministic fixture: academic metadata available, web backend unavailable.
        pf['capabilities']['academic.metadata']['available']=True; pf['capabilities']['academic.metadata']['status']='available'; pf['capabilities']['academic.metadata']['providers']=['fixture.academic']
        with tempfile.TemporaryDirectory() as td:
            f=Path(td)/'pf.json'; f.write_text(json.dumps(pf),encoding='utf-8')
            cp,d=run(ROOT/'scripts/router/resolve_bundle.py','Найди научные статьи по XRD','--route-id','academic-research','--stage','discovery','--preflight',f)
            self.assertEqual(cp.returncode,0,cp.stdout); self.assertEqual(d['bundle_state'],'READY'); self.assertEqual(d['route_id'],'academic-research'); self.assertIn(['web.discovery','academic.metadata'],d['required_any_of'])

    def test_writer_bundle_forbids_web(self):
        cp,pf=run(ROOT/'scripts/router/capability_preflight.py','--no-network')
        with tempfile.TemporaryDirectory() as td:
            f=Path(td)/'pf.json'; f.write_text(json.dumps(pf),encoding='utf-8')
            cp,d=run(ROOT/'scripts/router/resolve_bundle.py','Напиши научный отчет с графиком','--route-id','writing-prose','--stage','draft','--preflight',f)
            self.assertEqual(cp.returncode,0,cp.stdout); self.assertIn('web.discovery',d['forbidden_capabilities']); self.assertNotIn('search.discovery',d['logical_tools']); self.assertIn('reporting',d['domains'])

    def test_code_review_forbids_edit(self):
        cp,pf=run(ROOT/'scripts/router/capability_preflight.py','--no-network')
        with tempfile.TemporaryDirectory() as td:
            f=Path(td)/'pf.json'; f.write_text(json.dumps(pf),encoding='utf-8')
            cp,d=run(ROOT/'scripts/router/resolve_bundle.py','Проведи ревью кода','--route-id','code-review','--preflight',f)
            self.assertEqual(cp.returncode,0,cp.stdout); self.assertIn('code.edit',d['forbidden_capabilities']); self.assertNotIn('code.change',d['logical_tools'])

    def test_capability_escalation_is_bounded(self):
        cp,pf=run(ROOT/'scripts/router/capability_preflight.py','--no-network'); self.assertEqual(cp.returncode,0)
        with tempfile.TemporaryDirectory() as td:
            td=Path(td); pfpath=td/'pf.json'; pfpath.write_text(json.dumps(pf),encoding='utf-8')
            cp,b=run(ROOT/'scripts/router/resolve_bundle.py','Проанализируй PDF','--route-id','academic-research','--stage','document-analysis','--preflight',pfpath); self.assertEqual(cp.returncode,0,cp.stdout)
            bp=td/'bundle.json'; bp.write_text(json.dumps(b),encoding='utf-8')
            cp,d=run(ROOT/'scripts/router/capability_escalation.py',bp,'document.ocr','--reason','no_text_layer','--preflight',pfpath); self.assertEqual(cp.returncode,0,cp.stdout); self.assertEqual(d['decision'],'GRANTED')
            cp,d=run(ROOT/'scripts/router/capability_escalation.py',bp,'code.edit','--reason','need_fix','--preflight',pfpath); self.assertNotEqual(cp.returncode,0); self.assertEqual(d['decision'],'DENIED')

    def test_science(self):
        cp,d=run(ROOT/'scripts/capsules/science_compute.py','expr','m*c**2','--var','m=2','--var','c=3'); self.assertEqual(cp.returncode,0); self.assertEqual(d['result']['value'],18.0)
        cp,d=run(ROOT/'scripts/capsules/science_compute.py','equivalent','2*(x+1)','2*x+2'); self.assertEqual(cp.returncode,0); self.assertTrue(d['result']['equivalent'])
        cp,d=run(ROOT/'scripts/capsules/science_compute.py','uncertainty','a*b','--value','a=2','--value','b=3','--u','a=0.1','--u','b=0.2'); self.assertEqual(cp.returncode,0); self.assertGreater(d['result']['combined_standard_uncertainty'],0)
        cp,d=run(ROOT/'scripts/capsules/science_compute.py','molar-mass','Fe2O3'); self.assertEqual(cp.returncode,0); self.assertAlmostEqual(d['result']['molar_mass_g_mol'],159.6882,places=3)

    def test_docx_xlsx_tiff_pdf(self):
        from docx import Document
        import openpyxl, tifffile, numpy as np, fitz
        with tempfile.TemporaryDirectory() as td:
            td=Path(td)
            doc=Document(); doc.add_paragraph('alpha beta'); table=doc.add_table(rows=1,cols=2); table.rows[0].cells[0].text='A'; table.rows[0].cells[1].text='B'; dp=td/'x.docx'; doc.save(dp)
            cp,d=run(ROOT/'scripts/capsules/document_inspect.py',dp); self.assertEqual(cp.returncode,0,cp.stdout); self.assertEqual(d['kind'],'docx'); self.assertTrue(d['segments']); self.assertTrue(d['tables'])
            wb=openpyxl.Workbook(); ws=wb.active; ws['A1']=2; ws['B1']=3; ws['C1']='=A1+B1'; xp=td/'x.xlsx'; wb.save(xp)
            cp,d=run(ROOT/'scripts/capsules/document_inspect.py',xp); self.assertEqual(cp.returncode,0,cp.stdout); self.assertEqual(d['kind'],'xlsx'); self.assertEqual(d['sheets'][0]['cells'][2]['formula'],'=A1+B1')
            tp=td/'x.tiff'; tifffile.imwrite(tp,np.zeros((8,8),dtype=np.uint8)); cp,d=run(ROOT/'scripts/capsules/scientific_image_inspect.py',tp); self.assertEqual(cp.returncode,0,cp.stdout); self.assertFalse(d['phase_identification_allowed']); self.assertEqual(len(d['frames']),1)
            pp=td/'x.pdf'; pdf=fitz.open(); pg=pdf.new_page(); pg.insert_text((72,72),'hello diffraction'); pdf.save(pp); pdf.close()
            cp,d=run(ROOT/'scripts/capsules/document_inspect.py',pp); self.assertEqual(cp.returncode,0,cp.stdout); self.assertIn('diffraction',d['segments'][0]['text'])

    def test_corpus_incremental_and_delete(self):
        with tempfile.TemporaryDirectory() as td:
            td=Path(td); src=td/'src'; src.mkdir(); a=src/'a.md'; a.write_text('alpha diffraction beta',encoding='utf-8'); db=td/'c.db'
            cp,d=run(ROOT/'scripts/capsules/corpus_fts.py','index',src,'--db',db); self.assertEqual(cp.returncode,0,cp.stdout); self.assertEqual(d['result']['added'],1)
            cp,d=run(ROOT/'scripts/capsules/corpus_fts.py','index',src,'--db',db); self.assertEqual(cp.returncode,0); self.assertEqual(d['result']['skipped'],1)
            cp,d=run(ROOT/'scripts/capsules/corpus_fts.py','search','diffraction','--db',db); self.assertEqual(cp.returncode,0,cp.stdout); self.assertEqual(len(d['result']),1); self.assertTrue(d['result'][0]['locator'])
            a.unlink(); cp,d=run(ROOT/'scripts/capsules/corpus_fts.py','index',src,'--db',db); self.assertEqual(d['result']['deleted'],1)
            cp,d=run(ROOT/'scripts/capsules/corpus_fts.py','search','diffraction','--db',db); self.assertEqual(len(d['result']),0)

    def test_search_gateway_offline_fixture(self):
        with tempfile.TemporaryDirectory() as td:
            td=Path(td); f=td/'s.json'; f.write_text(json.dumps({'results':[{'title':'A','url':'https://x/a','content':'one','engine':'e'},{'title':'A2','url':'https://x/a','content':'dup','engine':'e'},{'title':'B','url':'https://x/b','content':'two','engine':'e'}]}),encoding='utf-8')
            cp,d=run(ROOT/'scripts/capsules/search_gateway.py','test query','--offline-fixture',f); self.assertEqual(cp.returncode,0,cp.stdout); self.assertEqual(len(d['results']),2); self.assertTrue(all(x['discovery_only'] for x in d['results']))

    def test_source_resolve_offline(self):
        cp,d=run(ROOT/'scripts/capsules/source_resolve.py','https://doi.org/10.1000/XYZ.123','--offline'); self.assertEqual(cp.returncode,0); self.assertEqual(d['doi_candidate'],'10.1000/xyz.123')

    def test_plot_and_report(self):
        with tempfile.TemporaryDirectory() as td:
            td=Path(td); csvp=td/'d.csv'; csvp.write_text('x,y\n1,2\n2,4\n',encoding='utf-8'); png=td/'p.png'
            cp,d=run(ROOT/'scripts/capsules/plot_render.py',csvp,'--x','x','--y','y','--output',png); self.assertEqual(cp.returncode,0,cp.stdout); self.assertTrue(png.exists())
            man=td/'r.json'; man.write_text(json.dumps({'title':'T','sections':[{'title':'S','body':'Body','claim_refs':['C1']}]}),encoding='utf-8'); md=td/'r.md'
            cp,d=run(ROOT/'scripts/capsules/report_compose.py',man,'--output',md); self.assertEqual(cp.returncode,0); self.assertIn('C1',md.read_text())

    def test_job_timeout_resume_and_reconcile(self):
        with tempfile.TemporaryDirectory() as td:
            td=Path(td); st=td/'job.json'; contract=td/'contract.json'; contract.write_text(json.dumps({'task':'research X'}),encoding='utf-8'); ctl=ROOT/'scripts/jobs/job_ctl.py'
            cp,d=run(ctl,'--state',st,'init','JOB-1','--contract',contract,'--stages','discovery,evidence'); self.assertEqual(cp.returncode,0,cp.stdout)
            cp,d=run(ctl,'--state',st,'attempt-start','--agent','researcher','--stage','discovery','--external-task-id','ext-1'); aid=d['attempt_id']
            cp,d=run(ctl,'--state',st,'checkpoint',aid,'--stage','discovery','--checkpoint','CHK-1'); self.assertEqual(cp.returncode,0)
            cp,d=run(ctl,'--state',st,'attempt-finish',aid,'TIMED_OUT'); self.assertEqual(d['job_status'],'WAITING_RETRY')
            cp,d=run(ctl,'--state',st,'stage','discovery','COMPLETED'); self.assertEqual(cp.returncode,0)
            cp,d=run(ctl,'--state',st,'stage','evidence','COMPLETED'); self.assertEqual(cp.returncode,0)
            cp,d=run(ctl,'--state',st,'child-add','FC-1','--mode','required'); self.assertEqual(cp.returncode,0)
            cp,d=run(ctl,'--state',st,'gate','evidence','PASS','--required'); self.assertEqual(cp.returncode,0)
            cp,d=run(ctl,'--state',st,'complete'); self.assertEqual(cp.returncode,4); self.assertEqual(d['error'],'COMPLETION_REJECTED')
            cp,d=run(ctl,'--state',st,'child-finish','FC-1','COMPLETED'); self.assertEqual(cp.returncode,0)
            cp,d=run(ctl,'--state',st,'complete'); self.assertEqual(cp.returncode,0,cp.stdout); self.assertEqual(d['status'],'COMPLETED')

if __name__=='__main__': unittest.main()
